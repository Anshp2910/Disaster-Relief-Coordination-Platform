#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generate the GTU Daily Learning Log (DLL) Word document (.docx) for the
Disaster Relief Coordination Platform internship.

Layout:
  * Cover page  – university / college / internship details + info table
  * Index page  – table of all 20 days (day, weekday, topic)
  * Day pages   – one page per day with Date / Day / Topic, a pre-written
                  learning summary, ruled lines for personal notes and
                  signature row.
"""

import os
import sys
from datetime import date, timedelta

# Ensure UTF-8 output on Windows consoles
try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT_DIR = r"C:\Disaster Relief Coordination Platform"
OUT_FILE = os.path.join(OUT_DIR, "Daily_Learning_Log.docx")

# ── Student / company details (editable) ──────────────────────────────────────
UNIVERSITY     = "Gujarat Technological University"
COLLEGE_NAME   = "C. K. Pithawala College of Engineering and Technology, Surat"
DEPARTMENT     = "Information Technology"
STUDENT_NAME   = "Ansh Patel"
ENROLMENT      = ""                      # fill in
FACULTY_GUIDE  = ""                      # fill in
DOMAIN         = "Web Development (Full-Stack MERN)"
COMPANY_NAME   = "Metanoia Solutions"
MENTOR_NAME    = "Nehal Patel"
MENTOR_PHONE   = ""                      # fill in
DURATION       = "1 July 2026 – 20 July 2026"
PROJECT_TITLE  = "Disaster Relief Coordination Platform"

# Day 1 of the log is a Wednesday; all 20 dates are derived from this.
START_DATE = date(2026, 7, 1)

# ── Colors ────────────────────────────────────────────────────────────────────
NAVY    = RGBColor(0x00, 0x2B, 0x5E)
ACCENT  = RGBColor(0x00, 0x70, 0xC0)
BLACK   = RGBColor(0x00, 0x00, 0x00)
GREY    = RGBColor(0x60, 0x60, 0x60)
HDR_BG  = "D0E0F8"
ALT_BG  = "F2F6FC"
LINE_CL = "B0BEC5"


# ── Daily entries: (weekday, topic, learning bullets) ────────────────────────
DAYS = [
    ("Wednesday", "Project setup, environment configuration, Git initialization",
     ["Set up the Node.js development environment (VS Code, Node, npm, Git, MongoDB).",
      "Created the client / server project structure for the MERN stack.",
      "Initialized the Git repository and configured .gitignore.",
      "Pushed the initial codebase to GitHub with a feature-branch workflow."]),
    ("Thursday", "Requirement analysis, MongoDB schema design",
     ["Studied disaster relief coordination requirements and identified user roles (citizen, volunteer, NGO, government, admin).",
      "Designed the core MongoDB collections: Users, Requests, Resources, Incidents, Zones, SOS Alerts.",
      "Defined field validations, references and indexes (including GeoJSON 2dsphere indexes for location queries).",
      "Documented the database schema for the report."]),
    ("Friday", "Express server setup, route structure implementation",
     ["Bootstrapped the Express.js server with essential middleware (helmet, cors, json, logger).",
      "Organised the REST routes into modules (auth, requests, resources, incidents, zones, sos).",
      "Integrated MongoDB through Mongoose with connection retry logic.",
      "Tested the base endpoints using Postman."]),
    ("Saturday", "JWT authentication implementation (register/login)",
     ["Implemented the register and login API endpoints.",
      "Hashed passwords securely using bcrypt.",
      "Generated signed JWT tokens and persisted the login session on the client.",
      "Added a protected-route middleware to verify tokens."]),
    ("Sunday", "React frontend setup, basic routing configuration",
     ["Scaffolded the React frontend using Vite with TypeScript.",
      "Configured React Router with lazy-loaded pages.",
      "Structured the source tree into components, pages, hooks, context and utils.",
      "Set up global styles and the theming foundation."]),
    ("Monday", "User registration form implementation with validation",
     ["Built the registration form component with reusable form fields.",
      "Added client-side validation (required fields, email format, password strength).",
      "Connected the form to the /api/auth/register endpoint.",
      "Displayed inline validation errors and a success toast."]),
    ("Tuesday", "Login page design with social auth buttons",
     ["Designed the login page with a glassmorphism card UI.",
      "Implemented email / password login flow with JWT storage.",
      "Added social authentication buttons (Google / GitHub) using Passport OAuth2.",
      "Stored the token securely and redirected based on user role."]),
    ("Wednesday", "Middleware implementation (auth, validation, error handling)",
     ["Implemented JWT authentication middleware for protected routes.",
      "Added Joi schema validation middleware for request bodies.",
      "Built centralised error-handling and response utilities.",
      "Applied rate limiting and input sanitisation on sensitive endpoints."]),
    ("Thursday", "Request CRUD API implementation",
     ["Built REST endpoints for relief requests (create, read, update, delete).",
      "Added filtering by status, priority and disaster type.",
      "Implemented pagination and keyword search.",
      "Wrote and executed API tests for the request endpoints."]),
    ("Friday", "Request list component with filtering",
     ["Created the request list component with a responsive card / table layout.",
      "Added a filter bar for status, priority and disaster type.",
      "Integrated the list with the requests API including debounced search.",
      "Added loading skeletons, empty states and pagination."]),
    ("Saturday", "Request detail page with status tracking",
     ["Built the request detail page with full metadata and attachments.",
      "Added a stepped progress tracker for request status.",
      "Implemented comments and activity feed for the request.",
      "Added role-based actions (assign, escalate, close)."]),
    ("Sunday", "Create request form with validation",
     ["Developed a multi-step create-request form.",
      "Added field-level validation and error messages.",
      "Integrated image / file upload with Multer.",
      "Wired submission to the API with a success redirect."]),
    ("Monday", "Resource model and API implementation",
     ["Designed the Resource model with categories and stock quantities.",
      "Implemented resource CRUD endpoints.",
      "Added allocation logic linking resources to requests.",
      "Added CSV bulk import / export utilities."]),
    ("Tuesday", "Resource management page with inventory display",
     ["Created the resource management page with an inventory table.",
      "Added quantity badges, filters and search.",
      "Implemented stock allocation and deallocation flows.",
      "Added empty / low-stock visual indicators."]),
    ("Wednesday", "Incident model and reporting system",
     ["Designed the Incident model with severity and disaster type fields.",
      "Built the incident reporting API with GeoJSON location support.",
      "Added incident status transitions and timestamps.",
      "Validated input with Joi and tested the endpoints."]),
    ("Thursday", "Incident list with filtering by severity and type",
     ["Created the incident list component with severity badges.",
      "Added filtering by severity and disaster type.",
      "Integrated incidents with the Leaflet map markers.",
      "Applied colour-coded severity styling and legends."]),
    ("Friday", "Dashboard page with KPI cards and recent requests",
     ["Built the dashboard with animated KPI cards (open requests, incidents, resources).",
      "Added the recent requests list and activity feed.",
      "Included charts for request trends by category.",
      "Made the dashboard role-aware for each user type."]),
    ("Saturday", "Role-based access control implementation",
     ["Defined the five user roles and their permissions.",
      "Implemented role checks in the API middleware.",
      "Guarded frontend routes per role in the router.",
      "Added conditional UI rendering based on the logged-in role."]),
    ("Sunday", "Admin dashboard with user management",
     ["Built the admin dashboard with statistics and charts.",
      "Added user management (list, search, activate / deactivate).",
      "Implemented escalation workflows for pending requests.",
      "Added data export and system configuration panels."]),
    ("Monday", "Admin user role management interface",
     ["Created the admin interface for managing user roles.",
      "Added role assignment and change functionality.",
      "Logged role changes for audit purposes.",
      "Performed final integration testing and updated documentation."]),
]


# ── Low-level helpers ─────────────────────────────────────────────────────────

def set_margins(doc):
    for sec in doc.sections:
        sec.top_margin = Inches(0.9)
        sec.bottom_margin = Inches(0.9)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)


def style_run(run, size=11, bold=False, italic=False, color=BLACK, name="Calibri"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return run


def para(doc, text="", size=11, bold=False, italic=False, color=BLACK,
         align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=6, line=1.15):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if text:
        style_run(p.add_run(text), size=size, bold=bold, italic=italic, color=color)
    return p


def heading(doc, text, size=13, color=NAVY, before=10, after=6):
    p = para(doc, text, size=size, bold=True, color=color, before=before, after=after)
    # accent underline
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "0070C0")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def bullet(doc, text, size=11, color=BLACK):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    style_run(p.add_run("•  "), size=size, bold=True, color=ACCENT)
    style_run(p.add_run(text), size=size, color=color)
    return p


def shade_cell(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def cell_text(cell, text, size=11, bold=False, italic=False, color=BLACK,
              align=WD_ALIGN_PARAGRAPH.LEFT, name="Calibri"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    style_run(p.add_run(text), size=size, bold=bold, italic=italic, color=color, name=name)
    return cell


def add_page_number_footer(doc, label="Daily Learning Log"):
    for sec in doc.sections:
        sec.footer.is_linked_to_previous = False
        p = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
        p.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p.add_run(f"{label}  |  Page ")
        style_run(r1, size=9, color=GREY)
        # PAGE field
        run = p.add_run()
        style_run(run, size=9, color=GREY)
        fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
        fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
        run._r.append(fld1); run._r.append(instr); run._r.append(fld2)
        r3 = p.add_run(" of ")
        style_run(r3, size=9, color=GREY)
        run2 = p.add_run()
        style_run(run2, size=9, color=GREY)
        f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = " NUMPAGES "
        f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
        run2._r.append(f1); run2._r.append(it); run2._r.append(f2)


def page_break(doc):
    doc.add_page_break()


def fmt_date(d):
    """Format a date as e.g. '1 July 2026'."""
    return f"{d.day} {d.strftime('%B')} {d.year}"


def ruled_lines(doc, count=8, height=2.0):
    """Empty paragraphs with a bottom border to simulate handwriting lines."""
    for _ in range(count):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = height
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), LINE_CL)
        pBdr.append(bottom)
        pPr.append(pBdr)


def label_value_table(doc, rows, col_widths=(Inches(2.3), Inches(4.4))):
    """Two-column label / value table used on the cover page."""
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, (label, value) in enumerate(rows):
        c0, c1 = table.rows[i].cells
        c0.width = col_widths[0]
        c1.width = col_widths[1]
        shade_cell(c0, HDR_BG)
        cell_text(c0, label, size=11, bold=True, color=NAVY)
        cell_text(c1, value if value else " ", size=11, color=BLACK)
    return table


# ── Build document ────────────────────────────────────────────────────────────

def build():
    doc = Document()
    set_margins(doc)
    add_page_number_footer(doc)

    # =====================================================================
    # COVER PAGE
    # =====================================================================
    para(doc, "", after=6)
    p = para(doc, UNIVERSITY, size=18, bold=True, color=NAVY,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    para(doc, COLLEGE_NAME, size=13, bold=True, color=ACCENT,
         align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    para(doc, f"Department of {DEPARTMENT}", size=11, color=GREY,
         align=WD_ALIGN_PARAGRAPH.CENTER, after=10)

    para(doc, "DAILY LEARNING LOG", size=28, bold=True, color=NAVY,
         align=WD_ALIGN_PARAGRAPH.CENTER, before=14, after=2)
    para(doc, "(Internship – Day-wise Learning Record)", size=12, italic=True,
         color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)

    para(doc, PROJECT_TITLE, size=15, bold=True, color=BLACK,
         align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    para(doc, f"Daily Log Period: {DURATION}", size=11, color=GREY,
         align=WD_ALIGN_PARAGRAPH.CENTER, after=16)

    label_value_table(doc, [
        ("Enrolment Number", ENROLMENT),
        ("Name of Students", STUDENT_NAME),
        ("Internal Faculty Guide", FACULTY_GUIDE),
        ("Domain", DOMAIN),
        ("Company’s Name", COMPANY_NAME),
        ("Company’s Mentor Name & Contact Number",
         (MENTOR_NAME + ("  /  " + MENTOR_PHONE if MENTOR_PHONE else "  /  ____________"))),
        ("Project Title", PROJECT_TITLE),
    ])

    para(doc, "", after=8)
    para(doc, "Note: Each day’s topic, learning summary and notes are recorded on the "
              "following pages. Important information may be underlined or highlighted "
              "in different colours for emphasis.",
         size=10, italic=True, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER)

    page_break(doc)

    # =====================================================================
    # INDEX PAGE
    # =====================================================================
    heading(doc, "Daily Schedule – Index", size=15, before=0)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(0.55), Inches(1.05), Inches(1.15), Inches(3.2), Inches(0.55)]
    for ci, w in enumerate(widths):
        table.columns[ci].width = w
    hdr = table.rows[0].cells
    for i, h in enumerate(["Day", "Date", "Weekday", "Topic", "Page"]):
        cell_text(hdr[i], h, size=11, bold=True, color=NAVY,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        shade_cell(hdr[i], HDR_BG)
    for idx, (weekday, topic, _bullets) in enumerate(DAYS, start=1):
        row = table.add_row().cells
        day_date = START_DATE + timedelta(days=idx - 1)
        cell_text(row[0], str(idx), align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(row[1], fmt_date(day_date), align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(row[2], weekday, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(row[3], topic)
        # Cover = page 1, index = page 2, so day N starts on page N + 2
        cell_text(row[4], str(idx + 2), align=WD_ALIGN_PARAGRAPH.CENTER)
        if idx % 2 == 0:
            for c in row:
                shade_cell(c, ALT_BG)

    page_break(doc)

    # =====================================================================
    # DAY PAGES
    # =====================================================================
    for idx, (weekday, topic, bullets) in enumerate(DAYS, start=1):
        # Day banner
        p = para(doc, f"DAILY LEARNING LOG", size=14, bold=True, color=NAVY,
                 align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=2)
        para(doc, f"Day {idx:02d}  •  {weekday}", size=11, bold=True, color=ACCENT,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=8)

        # Date / Day / Topic info table
        t = doc.add_table(rows=2, cols=4)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        r0 = t.rows[0].cells
        r1 = t.rows[1].cells
        for ci, w in enumerate([Inches(0.8), Inches(1.8), Inches(0.8), Inches(2.3)]):
            t.columns[ci].width = w
        day_date = START_DATE + timedelta(days=idx - 1)
        shade_cell(r0[0], HDR_BG); shade_cell(r0[2], HDR_BG)
        cell_text(r0[0], "Date", bold=True, color=NAVY)
        cell_text(r0[1], fmt_date(day_date))
        cell_text(r0[2], "Day", bold=True, color=NAVY)
        cell_text(r0[3], f"{idx} ({weekday})")
        # merged topic row
        merged = r1[0].merge(r1[1]).merge(r1[2]).merge(r1[3])
        shade_cell(merged, ALT_BG)
        cell_text(merged, f"Topic:  {topic}", bold=True, color=BLACK)

        # Learning summary
        heading(doc, "Learning Summary / Activities Performed", size=12)
        for b in bullets:
            bullet(doc, b)

        # Notes with ruled lines
        heading(doc, "Notes / Observations", size=12)
        para(doc, "You can underline and highlight important information with different "
                  "colours to emphasise it.", size=10, italic=True, color=GREY, after=6)
        ruled_lines(doc, count=7)

        # Signature row
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(0)
        tab_run = p.add_run("Intern Signature:  ______________________")
        style_run(tab_run, size=10, color=GREY)
        # Right-aligned tab just inside the 6.5" text area keeps the mentor
        # signature on the same line without overflowing the margin.
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.1),
                                                  WD_ALIGN_PARAGRAPH.RIGHT)
        tab_run = p.add_run("\tCompany Mentor Signature:  ______________________")
        style_run(tab_run, size=10, color=GREY)

        if idx < len(DAYS):
            page_break(doc)

    doc.save(OUT_FILE)
    print(f"Saved: {OUT_FILE}")


if __name__ == "__main__":
    build()
